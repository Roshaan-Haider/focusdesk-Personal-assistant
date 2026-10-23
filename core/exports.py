"""Export utilities for outreach and productivity data."""
from datetime import datetime
from pathlib import Path
import pandas as pd

from core.db import BASE_DIR

EXPORTS_DIR = BASE_DIR / "exports"


def export_outreach_to_file(rows: list[dict], fmt: str = "csv") -> Path:
    """Export outreach rows to a timestamped CSV or XLSX file.
    
    Columns: Date, Company, City, Recipient, Sender, Subject, Status, Campaign,
    Niche, Priority, Quoted Amount, Currency, Follow-up Date, Notes.
    """
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    ext = "xlsx" if fmt.lower() in ("xlsx", "excel") else "csv"
    filename = f"focusdesk_outreach_{ts}.{ext}"
    filepath = EXPORTS_DIR / filename

    data = []
    for r in rows:
        data.append({
            "Date": str(r.get("date_sent") or "")[:19],
            "Company": r.get("company_name") or "",
            "City": r.get("city") or "",
            "Recipient": r.get("recipient_email") or "",
            "Sender": r.get("sender_email") or "",
            "Subject": r.get("subject") or "",
            "Status": r.get("status") or "",
            "Campaign": r.get("campaign") or "",
            "Niche": r.get("tags") or r.get("lead_source") or "",
            "Priority": r.get("priority") or "",
            "Quoted Amount": r.get("quoted_amount") if r.get("quoted_amount") is not None else "",
            "Currency": r.get("currency") or "USD",
            "Follow-up Date": str(r.get("follow_up_date") or "")[:10],
            "Notes": r.get("notes") or "",
        })

    df = pd.DataFrame(data)

    if ext == "csv":
        df.to_csv(filepath, index=False, encoding="utf-8-sig")
    else:
        with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Outreach")
            ws = writer.sheets["Outreach"]
            ws.freeze_panes = "A2"

            from openpyxl.styles import Font, PatternFill
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="111A2E", end_color="111A2E", fill_type="solid")

            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill

            for col in ws.columns:
                col_letter = col[0].column_letter
                max_len = max(len(str(cell.value or "")) for cell in col)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    return filepath
